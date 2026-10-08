"""P2 — 특이일 AVG 제외 옵션, 분모 0 → N/A, 음수 수량 거부 (설계안 5장 추가 집계 기준)."""
import datetime as dt
import shutil
import subprocess
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook

from das_agent.engine import Scenario, compute, compare
from das_agent.ingestion import normalize
from das_agent.reporting import export_excel, explain_comparison
from das_agent.storage import Storage
from das_agent.tools import Toolbox

SOFFICE = next((p for p in [
    r"C:\Program Files\LibreOffice\program\soffice.exe",
    r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
    shutil.which("soffice") or "",
] if p and Path(p).exists()), None)


def three_day_frame() -> pd.DataFrame:
    """1/2 정상(주문 3), 1/3 주문 1건(특이일), 1/4 PCS 전부 0(토트 0 → 비율 N/A)."""
    rows = []
    d1, d2, d3 = dt.date(2025, 1, 2), dt.date(2025, 1, 3), dt.date(2025, 1, 4)
    rows += [(d1, "A 101", "X", 1, 10), (d1, "A 102", "X", 2, 20), (d1, "A 102", "Y", 1, 1), (d1, "B 201", "X", 1, 16)]
    rows += [(d2, "C 301", "X", 3, 45)]
    rows += [(d3, "D 401", "X", 1, 0), (d3, "D 402", "Y", 1, 0)]
    return pd.DataFrame(rows, columns=["date", "order_key", "sku", "order_line", "pcs"])


@pytest.fixture
def res():
    return compute(three_day_frame(), Scenario("기준", 2, 15))


def test_outlier_days_flagged_and_zero_denominator_is_na(res):
    assert res.flags["single_order_days"] == ["2025-01-03"]
    assert res.flags["outlier_days"] == ["2025-01-03"]        # 1/4는 비율 N/A라 OL/TOTE<1 아님
    day3 = res.daily[res.daily["date"] == "2025-01-04"].iloc[0]
    assert day3["tote"] == 0 and pd.isna(day3["ol_per_tote"]) and pd.isna(day3["pcs_per_tote"])


def test_avg_excludes_zero_and_avg_ex_excludes_outliers_only(res):
    d = res.daily.set_index(res.daily["date"].dt.strftime("%Y-%m-%d"))
    # AVG: 0·빈 값 제외 → 1/2, 1/3 (1/4는 N/A)
    assert res.avg["ol_per_tote"] == pytest.approx((d.loc["2025-01-02", "ol_per_tote"] + d.loc["2025-01-03", "ol_per_tote"]) / 2)
    # AVG(특이일 제외): 1/3 제외 → 1/2 만
    assert res.avg_ex["ol_per_tote"] == pytest.approx(d.loc["2025-01-02", "ol_per_tote"])
    assert res.avg_ex["days_included"] == 2 and res.avg_ex["days_excluded"] == 1
    # 합계·기간 비율은 특이일 포함 그대로
    assert res.period["order_line"] == 10 and res.period["pcs"] == 92
    assert res.period["ol_per_tote"] == pytest.approx(10 / res.period["tote"])


def test_compare_table_and_explanation_show_outlier_counts(res):
    other = compute(three_day_frame(), Scenario("3셀", 3, 15))
    table = compare([res, other])
    assert list(table["특이일 수"]) == [1, 1]
    assert table.loc[0, "OL/TOTE(일평균·특이일 제외)"] == pytest.approx(res.avg_ex["ol_per_tote"])
    text = explain_comparison([res, other])
    assert "특이일 1일" in text and "포함 2일 / 제외 1일" in text


def test_storage_roundtrip_and_tool_output_keep_avg_ex(tmp_path):
    st = Storage(tmp_path / "db")
    frame = three_day_frame()
    frame["src_row"] = range(2, 2 + len(frame))
    frame["category"] = "미분류"
    ds = st.register_dataset(frame, name="t", file_name="t.csv", file_hash="t", sheet="csv", mapping={}, validation={})
    rid, r1 = st.run_scenario(ds, Scenario("기준", 2, 15))
    loaded = st.load_run(rid)
    assert loaded.avg_ex == r1.avg_ex and loaded.flags["outlier_days"] == ["2025-01-03"]
    out = Toolbox(st).simulate_scenarios(ds, [{"name": "기준", "cells": 2, "tote_capacity": 15}])
    ex = out["runs"][0]["avg_excl_outliers"]
    assert ex["days_included"] == 2 and ex["ol_per_tote"] == pytest.approx(r1.avg_ex["ol_per_tote"])
    st.close()


def test_export_has_outlier_column_and_second_avg_row(res, tmp_path):
    out = export_excel([res], tmp_path / "o.xlsx")
    wb = load_workbook(out)
    ws = wb[wb.sheetnames[1]]
    assert ws.cell(4, 16).value == "특이일"
    flags = {str(ws.cell(5 + i, 1).value)[:10]: ws.cell(5 + i, 16).value for i in range(res.n_days)}
    assert flags == {"2025-01-02": 0, "2025-01-03": 1, "2025-01-04": 0}
    tot = 5 + res.n_days
    assert ws.cell(tot + 2, 1).value.startswith("AVG(특이일 제외")
    assert "AVERAGEIFS" in ws.cell(tot + 2, 8).value
    cached = load_workbook(out, data_only=True)[wb.sheetnames[1]]
    assert cached.cell(tot + 2, 8).value == pytest.approx(res.avg_ex["ol_per_tote"])
    assert cached.cell(tot, 8).value == pytest.approx(res.period["ol_per_tote"])
    cmp_ws = load_workbook(out, data_only=True)["비교"]
    assert cmp_ws.cell(3, 26).value == "특이일 수" and cmp_ws.cell(4, 26).value == 1


@pytest.mark.skipif(SOFFICE is None, reason="LibreOffice 없음")
def test_recalculated_avg_rows_match_engine(res, tmp_path):
    out = export_excel([res], tmp_path / "o.xlsx")
    rd = tmp_path / "recalc"
    rd.mkdir()
    subprocess.run([SOFFICE, "--headless", "--calc", "--convert-to", "xlsx", "--outdir", str(rd), str(out)],
                   check=True, capture_output=True, timeout=180)
    wb = load_workbook(rd / "o.xlsx", data_only=True)
    ws = wb[wb.sheetnames[1]]
    tot = 5 + res.n_days
    assert ws.cell(tot + 1, 8).value == pytest.approx(res.avg["ol_per_tote"])
    assert ws.cell(tot + 2, 8).value == pytest.approx(res.avg_ex["ol_per_tote"])
    assert ws.cell(tot + 2, 9).value == pytest.approx(res.avg_ex["pcs_per_tote"])
    assert ws.cell(tot, 16).value == 1
    day3 = next(i for i in range(5, tot) if str(ws.cell(i, 1).value)[:10] == "2025-01-04")
    assert ws.cell(day3, 8).value == "N/A"
    errors = [c.coordinate for s in wb.worksheets for row in s.iter_rows() for c in row
              if isinstance(c.value, str) and c.value.startswith("#")]
    assert not errors


def test_negative_and_decimal_quantities_are_errors_not_dropped():
    raw = pd.DataFrame({"실확정일자": ["20250102", "20250102"], "배송처차수": ["A 1", "A 2"],
                        "상품코드": ["X", "Y"], "Order Line": ["1", "1"], "PCS": ["-3", "2.5"]})
    df, val = normalize(raw, {"date": "실확정일자", "order_key": "배송처차수", "sku": "상품코드",
                              "order_line": "Order Line", "pcs": "PCS", "category": None})
    assert not val.ok
    assert len(val.errors) == 2 and all("PCS" in e["reason"] or "pcs" in e["reason"].lower() for e in val.errors)
