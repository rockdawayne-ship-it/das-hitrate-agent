"""엑셀 내보내기: 수식 오류 0개, 합계가 화면 KPI와 일치 (PRD 성공 기준 5)."""
import math
import shutil
import subprocess
from pathlib import Path

import pytest
from openpyxl import load_workbook

from das_agent.engine import Scenario, compute, compare
from das_agent.ingestion import make_sample_file, read_raw, normalize, suggest_mapping
from das_agent.reporting import export_excel, explain_comparison

SOFFICE = next((p for p in [
    r"C:\Program Files\LibreOffice\program\soffice.exe",
    r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
    shutil.which("soffice") or "",
] if p and Path(p).exists()), None)


@pytest.fixture(scope="module")
def results(tmp_path_factory):
    tmp = tmp_path_factory.mktemp("rep")
    xlsx = make_sample_file(tmp / "sample.xlsx", days=20)
    raw = read_raw(xlsx, "data")
    df, val = normalize(raw, suggest_mapping(list(raw.columns)))
    assert val.ok
    scs = [Scenario("16셀·15PCS", 16, 15), Scenario("24셀·20PCS", 24, 20), Scenario("24셀·20PCS 배송처", 24, 20, "배송처")]
    return [compute(df, s) for s in scs]


def test_export_structure_and_values(results, tmp_path):
    out = export_excel(results, tmp_path / "out.xlsx", dataset={"name": "sample", "file_name": "sample.xlsx",
                                                                    "sheet": "data", "row_count": 1, "mapping": {}})
    wb = load_workbook(out)
    assert wb.sheetnames[0] == "비교" and "설정" in wb.sheetnames and "계산정의" in wb.sheetnames
    assert len(wb.sheetnames) == 3 + len(results)
    ws = wb[wb.sheetnames[1]]
    r0 = results[0]
    n = r0.n_days
    tot = 5 + n
    assert ws.cell(tot, 1).value.startswith("합계")
    assert ws.cell(tot, 2).value == r0.period["order"]
    assert ws.cell(tot, 7).value == f"=SUM(G5:G{4 + n})"
    assert sum(ws.cell(5 + i, 7).value for i in range(n)) == r0.period["tote"]
    cmp_ws = wb["비교"]
    assert cmp_ws.cell(4, 7).value == "●"
    assert cmp_ws.cell(4, 12).value == f"='{wb.sheetnames[1]}'!G{tot}"


@pytest.mark.skipif(SOFFICE is None, reason="LibreOffice 없음: 수식 재계산 검증 생략")
def test_recalculated_formulas_match_engine(results, tmp_path):
    out = export_excel(results, tmp_path / "out.xlsx")
    recalc_dir = tmp_path / "recalc"
    recalc_dir.mkdir()
    subprocess.run([SOFFICE, "--headless", "--calc", "--convert-to", "xlsx", "--outdir", str(recalc_dir), str(out)],
                   check=True, capture_output=True, timeout=180)
    wb = load_workbook(recalc_dir / "out.xlsx", data_only=True)
    errors = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("#"):
                    errors.append((ws.title, c.coordinate, c.value))
    assert errors == [], f"수식 오류 셀: {errors[:10]}"
    cmp_df = compare(results)
    for i, r in enumerate(results):
        ws = wb[wb.sheetnames[1 + i]]
        tot = 5 + r.n_days
        assert ws.cell(tot, 7).value == r.period["tote"]
        assert ws.cell(tot, 8).value == pytest.approx(r.period["ol_per_tote"], rel=1e-9)
        assert ws.cell(tot, 9).value == pytest.approx(r.period["pcs_per_tote"], rel=1e-9)
        assert ws.cell(tot, 14).value == r.period["tote_per_h"]
        assert ws.cell(tot + 1, 8).value == pytest.approx(r.avg["ol_per_tote"], rel=1e-9)
        assert ws.cell(tot + 1, 7).value == pytest.approx(r.avg["tote"], rel=1e-9)
        c = wb["비교"]
        assert c.cell(4 + i, 12).value == r.period["tote"]
        assert c.cell(4 + i, 16).value == cmp_df.loc[i, "토트 증감"]
        assert c.cell(4 + i, 13).value == pytest.approx(r.period["ol_per_tote"], rel=1e-9)


def test_explain_comparison_text(results):
    txt = explain_comparison(results)
    assert "기준 '16셀·15PCS'" in txt
    assert f"총 토트 {results[1].period['tote']:,}" in txt
    assert "가장 적은" in txt or "동률" in txt
