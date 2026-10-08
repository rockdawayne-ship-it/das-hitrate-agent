import datetime as dt
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook

from das_agent.agent import Agent, AgentContext
from das_agent.engine import Scenario, Filters, golden_example, compute
from das_agent.ingestion import normalize, suggest_mapping, read_raw, inspect_file
from das_agent.llm import NullProvider, OllamaProvider, get_provider
from das_agent.reporting import export_excel, explain_comparison
from das_agent.storage import Storage
from das_agent.tools import Toolbox, ToolError


@pytest.fixture
def setup(tmp_path):
    storage = Storage(tmp_path / "db")
    frame = golden_example()
    frame["src_row"] = range(2, 6)
    dataset = storage.register_dataset(frame, name="PRD", file_name="g.csv", file_hash="g", sheet="csv", mapping={}, validation={})
    context = AgentContext(dataset_id=dataset, baseline=Scenario("기준", 16, 15), date_min="2025-01-02", date_max="2025-01-02")
    tools = Toolbox(storage)
    yield storage, tools, context, Agent(tools, NullProvider())
    storage.close()


def raw_data():
    return pd.DataFrame({"실확정일자": ["20250102"], "배송처차수": ["00123"], "상품코드": ["00007"], "Order Line": ["1"], "PCS": ["4"]})


@pytest.mark.parametrize("value", ["inf", "-inf", "1e50", "3.5", "-1", ""])
def test_bad_quantities_report_instead_of_crash(value):
    raw = raw_data()
    raw.loc[0, "PCS"] = value
    _, validation = normalize(raw, suggest_mapping(raw.columns))
    assert not validation.ok


def test_excel_preserves_text_identifiers_and_case_insensitive_sheet(tmp_path):
    raw = raw_data()
    raw.loc[0, "상품코드"] = "NA"
    path = tmp_path / "source.xlsx"
    raw.to_excel(path, sheet_name="DATA", index=False)
    inspection = inspect_file(path, "data")
    loaded = read_raw(path, inspection.sheet)
    frame, validation = normalize(loaded, inspection.mapping)
    assert validation.ok
    assert frame.order_key.iloc[0] == "00123"
    assert frame.sku.iloc[0] == "NA"
    assert inspection.sheet == "DATA"
    with pytest.raises(ValueError, match="시트"):
        read_raw(path, "missing")


def test_empty_and_duplicate_mapping_rejected():
    raw = raw_data()
    mapping = suggest_mapping(raw.columns)
    with pytest.raises(ValueError, match="데이터 행"):
        normalize(raw.iloc[:0], mapping)
    mapping["pcs"] = mapping["order_line"]
    with pytest.raises(ValueError, match="중복"):
        normalize(raw, mapping)


@pytest.mark.parametrize("value", ["202511", "2025012", "01/02/2025", "2025-02-30"])
def test_dates_need_explicit_year_month_day(value):
    raw = raw_data()
    raw.loc[0, "실확정일자"] = value
    _, validation = normalize(raw, suggest_mapping(raw.columns))
    assert not validation.ok


@pytest.mark.parametrize("key,value", [("cells", 2.5), ("tote_capacity", True), ("hours_per_day", float("inf")), ("hours_per_day", 25)])
def test_tool_strict_numeric_validation(setup, key, value):
    _, tools, ctx, _ = setup
    with pytest.raises(ToolError):
        tools.simulate_scenarios(ctx.dataset_id, [{"cells": 16, "tote_capacity": 15, key: value}])


def test_negative_decimal_requests_not_silently_rounded(setup):
    _, _, ctx, agent = setup
    for text in ("토트 -5로 계산", "2.5셀 비교", "토트 용량 3.5로 계산", "0시간으로 계산"):
        reply = agent.handle(text, ctx)
        assert reply.needs_input, text
        assert not reply.run_ids


def test_actual_baseline_is_selected_when_not_first_in_grid(setup):
    _, _, ctx, agent = setup
    ctx.baseline = Scenario("기준", 24, 20)
    reply = agent.handle("16·24·32셀 토트 15·20PCS 비교", ctx)
    baseline = reply.compare_table.query("기준 == '●'").iloc[0]
    assert baseline["DAS 셀"] == 24 and baseline["토트 용량"] == 20
    assert reply.run_ids.index(reply.baseline_run_id) == 3


def test_explicit_filter_clear_and_followup_grid(setup):
    _, _, ctx, agent = setup
    agent.handle("16·24·32셀 토트 15·20PCS 비교", ctx)
    agent.handle("2025-01-02 ~ 2025-01-02 기간으로 계산", ctx)
    assert ctx.filters.start == "2025-01-02"
    reply = agent.handle("전체 기간으로 다시 계산", ctx)
    assert ctx.filters == Filters()
    assert len(reply.run_ids) == 6
    reply = agent.handle("배송처로 다시 계산", ctx)
    assert len(reply.run_ids) == 7
    changed = reply.compare_table.query("`주문 단위` == '배송처'")
    assert {(r["DAS 셀"], r["토트 용량"]) for _, r in changed.iterrows()} == {(c, t) for c in (16, 24, 32) for t in (15, 20)}


def test_invalid_dates_return_clear_error(setup):
    _, _, ctx, agent = setup
    for text in ("2025-02-30 ~ 2025-03-01 계산", "2025-03-02 ~ 2025-01-01 계산", "2025-13월 계산"):
        reply = agent.handle(text, ctx)
        assert reply.needs_input, text


def test_zero_tote_comparison_and_cached_formulas(tmp_path):
    frame = golden_example()
    frame["pcs"] = 0
    results = [compute(frame, Scenario("기준", 2, 15)), compute(frame, Scenario("변경", 3, 20))]
    assert "N/A" in explain_comparison(results)
    output = export_excel(results, tmp_path / "zero.xlsx")
    workbook = load_workbook(output, data_only=True)
    assert workbook["비교"]["M4"].value == "N/A"
    assert workbook["비교"]["L4"].value == 0


def test_formula_cache_matches_engine_and_user_names_are_text(tmp_path):
    result = compute(golden_example(), Scenario("=1+2", 2, 15))
    out = export_excel([result], tmp_path / "out.xlsx", dataset={"name": "=1+2", "file_name": "=3+4"})
    cached = load_workbook(out, data_only=True)
    assert cached["비교"]["L4"].value == 5
    assert cached["비교"]["M4"].value == 1
    assert cached["비교"]["N4"].value == 9.4
    formulas = load_workbook(out)
    assert formulas["비교"]["A4"].value == "=1+2"
    assert formulas["비교"]["A4"].data_type == "s"
    assert formulas["설정"]["B3"].data_type == "s"


def test_cache_preserves_requested_name_after_reopen(setup):
    storage, _, ctx, _ = setup
    a, _ = storage.run_scenario(ctx.dataset_id, Scenario("원래 이름"))
    b, _ = storage.run_scenario(ctx.dataset_id, Scenario("새 이름"))
    assert storage.load_run(b).scenario.name == "새 이름"
    assert storage.load_run(a).scenario.name == "원래 이름"


def test_export_rejects_mixed_dataset(setup):
    storage, tools, ctx, _ = setup
    frame = golden_example()
    frame["src_row"] = range(2, 6)
    other = storage.register_dataset(frame, name="other", file_name="other", file_hash="o", sheet="csv", mapping={}, validation={})
    a, _ = storage.run_scenario(ctx.dataset_id, Scenario())
    b, _ = storage.run_scenario(other, Scenario())
    with pytest.raises(ToolError, match="다른 데이터셋"):
        tools.export_report([a, b])


def test_no_cloud_or_remote_fallback(monkeypatch):
    monkeypatch.setenv("DAS_LLM", "anthropic")
    assert isinstance(get_provider(), NullProvider)
    with pytest.raises(ValueError):
        OllamaProvider(host="https://example.com")
    with pytest.raises(ValueError):
        OllamaProvider(model="some-model:cloud")
