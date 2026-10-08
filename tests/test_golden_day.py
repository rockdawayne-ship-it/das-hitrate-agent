"""PRD 성공 기준 2 — 1일치 정답 대조 (합성 데이터).

실데이터 수작업 정답이 확보되기 전까지, 같은 절차를 엔진과 독립된 두 경로로 재현해 모든 열을 대조한다.
  1. 손 절차(순수 파이썬): PRD 4장을 pandas 없이 dict·sorted()로 그대로 따라 한다.
  2. 엑셀 수식: MATCH 기반 웨이브 번호 + SUMIFS 피벗 + ROUNDUP, LibreOffice로 재계산한 값.
세 경로의 날짜별·기간 값이 모든 열에서 일치해야 한다 (주문 단위 2종).
"""
import math
import shutil
import subprocess
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import Workbook, load_workbook

from das_agent.engine import Scenario, compute
from das_agent.ingestion import make_golden_day_file, normalize, read_raw, suggest_mapping

SOFFICE = next((p for p in [
    r"C:\Program Files\LibreOffice\program\soffice.exe",
    r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
    shutil.which("soffice") or "",
] if p and Path(p).exists()), None)

CELLS, CAP, HOURS = 16, 15, 10.0
COLS = ["order", "sku", "order_line", "pcs", "waves", "tote", "ol_per_tote", "pcs_per_tote",
        "ol_per_order", "pcs_per_order", "pcs_per_ol", "ol_per_sku", "tote_per_h", "pcs_per_h"]


def unit_key(order_key: str, unit: str) -> str:
    return order_key if unit == "배송처차수" else order_key.split(" ", 1)[0]


def hand_procedure(rows: list[dict], cells: int, cap: int, unit: str, hours: float) -> dict:
    """PRD 4장을 손으로 하듯 따라 한 계산. 엔진 코드를 전혀 쓰지 않는다."""
    keys = sorted({unit_key(r["배송처차수"], unit) for r in rows})       # 문자 오름차순
    wave_of = {k: i // cells + 1 for i, k in enumerate(keys)}            # 셀 수만큼 묶음
    pivot: dict[tuple[int, str], int] = {}
    for r in rows:
        w = wave_of[unit_key(r["배송처차수"], unit)]
        pivot[(w, r["상품코드"])] = pivot.get((w, r["상품코드"]), 0) + int(r["PCS"])
    tote = sum(math.ceil(v / cap) for v in pivot.values())               # 합산 후 올림
    order, sku = len(keys), len({r["상품코드"] for r in rows})
    ol, pcs = sum(int(r["Order Line"]) for r in rows), sum(int(r["PCS"]) for r in rows)
    return {"order": order, "sku": sku, "order_line": ol, "pcs": pcs, "waves": max(wave_of.values()), "tote": tote,
            "ol_per_tote": ol / tote, "pcs_per_tote": pcs / tote, "ol_per_order": ol / order,
            "pcs_per_order": pcs / order, "pcs_per_ol": pcs / ol, "ol_per_sku": ol / sku,
            "tote_per_h": math.ceil(tote / hours), "pcs_per_h": pcs / hours}


def build_manual_workbook(rows: list[dict], cells: int, cap: int, unit: str, hours: float, path: Path) -> Path:
    """엑셀 수작업(피벗+ROUNDUP)을 수식으로 재현한 통합문서. 값은 LibreOffice 재계산으로 얻는다."""
    wb = Workbook()
    ws = wb.active
    ws.title = "data"
    ws.append(["실확정일자", "배송처차수", "상품코드", "Order Line", "PCS", "주문키", "웨이브"])
    n = len(rows)
    keys = sorted({unit_key(r["배송처차수"], unit) for r in rows})
    for i, r in enumerate(rows, start=2):
        ws.append([int(r["실확정일자"]), str(r["배송처차수"]), str(r["상품코드"]), int(r["Order Line"]), int(r["PCS"])])
        ws.cell(i, 6, f"=B{i}" if unit == "배송처차수" else f'=LEFT(B{i},FIND(" ",B{i})-1)')
        ws.cell(i, 7, f"=ROUNDUP(MATCH(F{i},keys!$A$1:$A${len(keys)},0)/{cells},0)")
        ws.cell(i, 2).data_type = "s"
        ws.cell(i, 3).data_type = "s"
    wk = wb.create_sheet("keys")
    for i, k in enumerate(keys, start=1):
        wk.cell(i, 1, k).data_type = "s"
    # 피벗: (웨이브, SKU) 조합 목록은 파이썬이 나열하지만, 합계·올림은 엑셀 수식이 계산한다.
    wave_of = {k: i // cells + 1 for i, k in enumerate(keys)}
    pairs = sorted({(wave_of[unit_key(r["배송처차수"], unit)], r["상품코드"]) for r in rows})
    wp = wb.create_sheet("pivot")
    wp.append(["웨이브", "상품코드", "PCS합", "토트"])
    for i, (w, s) in enumerate(pairs, start=2):
        wp.cell(i, 1, w)
        wp.cell(i, 2, s).data_type = "s"
        wp.cell(i, 3, f"=SUMIFS(data!$E$2:$E${n + 1},data!$G$2:$G${n + 1},A{i},data!$C$2:$C${n + 1},B{i})")
        wp.cell(i, 4, f"=ROUNDUP(C{i}/{cap},0)")
    m = len(pairs) + 1
    wsum = wb.create_sheet("summary")
    formulas = {
        "order": f"=COUNTA(keys!A1:A{len(keys)})",
        "sku": f"=SUMPRODUCT(1/COUNTIF(data!C2:C{n + 1},data!C2:C{n + 1}))",
        "order_line": f"=SUM(data!D2:D{n + 1})",
        "pcs": f"=SUM(data!E2:E{n + 1})",
        "waves": f"=MAX(data!G2:G{n + 1})",
        "tote": f"=SUM(pivot!D2:D{m})",
        "pivot_pcs": f"=SUM(pivot!C2:C{m})",          # 피벗 PCS 합 = 원본 PCS 합이어야 함
        "ol_per_tote": "=B3/B6", "pcs_per_tote": "=B4/B6", "ol_per_order": "=B3/B1",
        "pcs_per_order": "=B4/B1", "pcs_per_ol": "=B4/B3", "ol_per_sku": "=B3/B2",
        "tote_per_h": f"=ROUNDUP(B6/{hours},0)", "pcs_per_h": f"=B4/{hours}",
    }
    for i, (k, f) in enumerate(formulas.items(), start=1):
        wsum.cell(i, 1, k)
        wsum.cell(i, 2, f)
    wb.save(path)
    return path


def recalc_with_libreoffice(path: Path) -> Path:
    out_dir = path.parent / "recalc"
    out_dir.mkdir(exist_ok=True)
    subprocess.run([SOFFICE, "--headless", "--calc", "--convert-to", "xlsx", "--outdir", str(out_dir), str(path)],
                   check=True, capture_output=True, timeout=180)
    return out_dir / path.name


def read_summary(path: Path) -> dict:
    ws = load_workbook(path, data_only=True)["summary"]
    return {ws.cell(i, 1).value: ws.cell(i, 2).value for i in range(1, ws.max_row + 1)}


def close(a, b) -> bool:
    return math.isclose(float(a), float(b), rel_tol=0, abs_tol=1e-9)


@pytest.fixture(scope="module")
def golden(tmp_path_factory):
    tmp = tmp_path_factory.mktemp("golden")
    xlsx = make_golden_day_file(tmp / "golden_day.xlsx")
    raw = read_raw(xlsx, "data")
    df, val = normalize(raw, suggest_mapping(list(raw.columns)))
    assert val.ok, val.to_dict()
    rows = raw.to_dict("records")
    return tmp, rows, df


def test_golden_day_has_intended_edge_cases(golden):
    _, rows, _ = golden
    keys = sorted({r["배송처차수"] for r in rows})
    assert keys.index("10000 301") < keys.index("9000 301")        # 문자순 (숫자순 아님)
    assert keys.index("A10 301") < keys.index("A2 301")
    assert len(keys) == 50 and len({k.split(" ")[0] for k in keys}) == 49   # 배송처 묶으면 49
    assert sum(1 for r in rows if r["배송처차수"] == "A1 301") == 2          # 분할 행


@pytest.mark.parametrize("unit", ["배송처차수", "배송처"])
def test_engine_matches_hand_procedure_all_columns(golden, unit):
    _, rows, df = golden
    res = compute(df, Scenario("정답", CELLS, CAP, unit, HOURS))
    assert res.n_days == 1
    daily = res.daily.iloc[0].to_dict()
    hand = hand_procedure(rows, CELLS, CAP, unit, HOURS)
    mismatch = {c: (daily[c], hand[c]) for c in COLS if not close(daily[c], hand[c])}
    assert not mismatch, mismatch
    # 하루짜리 데이터는 기간 합계 = 날짜별 값 (tote/h 정의 포함)
    for c in COLS:
        assert close(res.period[c], hand[c]), (c, res.period[c], hand[c])
    assert res.period["order"] == (50 if unit == "배송처차수" else 49)


@pytest.mark.skipif(SOFFICE is None, reason="LibreOffice 없음: 엑셀 수식 대조 생략")
@pytest.mark.parametrize("unit", ["배송처차수", "배송처"])
def test_engine_matches_excel_manual_formulas(golden, unit):
    tmp, rows, df = golden
    book = build_manual_workbook(rows, CELLS, CAP, unit, HOURS, tmp / f"manual_{unit}.xlsx")
    vals = read_summary(recalc_with_libreoffice(book))
    assert all(not (isinstance(v, str) and v.startswith("#")) for v in vals.values()), vals
    assert vals["pivot_pcs"] == vals["pcs"], "피벗 PCS 합이 원본과 다름 → 웨이브 배정 불일치"
    res = compute(df, Scenario("정답", CELLS, CAP, unit, HOURS))
    daily = res.daily.iloc[0].to_dict()
    mismatch = {c: (daily[c], vals[c]) for c in COLS if not close(daily[c], vals[c])}
    assert not mismatch, mismatch
    hand = hand_procedure(rows, CELLS, CAP, unit, HOURS)
    assert all(close(vals[c], hand[c]) for c in COLS)
