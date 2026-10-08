"""비교표·근거 설명·엑셀 내보내기 (PRD 5장 '엑셀 내보내기', 설계안 7장)."""
from __future__ import annotations

import re
import io
import math
import zipfile
import xml.etree.ElementTree as ET
from datetime import datetime, timezone, timedelta
from pathlib import Path
from typing import Optional

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from . import ENGINE_VERSION
from .engine import RunResult, METRIC_DEFS, compare

HEAD_FILL = PatternFill("solid", fgColor="DDEBF7")
SUM_FILL = PatternFill("solid", fgColor="FFF2CC")
BOLD = Font(bold=True)


def _sheet_name(i: int, r: RunResult) -> str:
    base = f"S{i + 1}_{r.scenario.cells}c_{r.scenario.tote_capacity}t_{'차수' if r.scenario.order_unit == '배송처차수' else '배송처'}"
    return re.sub(r"[\[\]\*\?/\\:]", "_", base)[:31]


def _write_header(ws, row: int, headers: list[str]) -> None:
    for j, h in enumerate(headers, start=1):
        c = ws.cell(row=row, column=j, value=h)
        c.font = BOLD
        c.fill = HEAD_FILL
        c.alignment = Alignment(horizontal="center")


def _autosize(ws, widths: Optional[dict] = None) -> None:
    for col in ws.columns:
        letter = get_column_letter(col[0].column)
        w = max((len(str(c.value)) for c in col if c.value is not None), default=8)
        ws.column_dimensions[letter].width = min(max(10, w + 2), 40)
    for k, v in (widths or {}).items():
        ws.column_dimensions[k].width = v


def write_scenario_sheet(wb: Workbook, title: str, r: RunResult) -> None:
    """날짜별 결과 시트. 합계·비율·시간당은 엑셀 수식."""
    ws = wb.create_sheet(title)
    ws["A1"] = f"시나리오: {r.scenario.name}"
    ws["A1"].font = BOLD
    ws["A2"] = "DAS 셀"; ws["B2"] = r.scenario.cells
    ws["C2"] = "토트 용량"; ws["D2"] = r.scenario.tote_capacity
    ws["E2"] = "주문 단위"; ws["F2"] = r.scenario.order_unit
    ws["G2"] = "작업시간"; ws["H2"] = r.scenario.hours_per_day   # $H$2 참조
    ws["I2"] = "필터"; ws["J2"] = r.filters.describe()
    headers = ["날짜", "order", "sku", "order line", "pcs", "waves", "tote",
               "HIT RATE OL/TOTE", "HIT RATE PCS/TOTE", "주문당 OL", "주문당 PCS", "OL당 PCS", "SKU당 OL",
               "tote/h", "PCS/H", "특이일"]
    HR = 4
    _write_header(ws, HR, headers)
    first = HR + 1
    daily = r.daily
    outlier_set = set(r.flags.get("outlier_days", []))
    for i, row in enumerate(daily.itertuples(index=False)):
        rr = first + i
        ws.cell(rr, 1, pd.Timestamp(row.date).date()).number_format = "yyyy-mm-dd"
        ws.cell(rr, 2, int(row.order)); ws.cell(rr, 3, int(row.sku))
        ws.cell(rr, 4, int(row.order_line)); ws.cell(rr, 5, int(row.pcs))
        ws.cell(rr, 6, int(row.waves)); ws.cell(rr, 7, int(row.tote))
        ws.cell(rr, 8, f'=IF(G{rr}=0,"N/A",D{rr}/G{rr})')
        ws.cell(rr, 9, f'=IF(G{rr}=0,"N/A",E{rr}/G{rr})')
        ws.cell(rr, 10, f'=IF(B{rr}=0,"N/A",D{rr}/B{rr})')
        ws.cell(rr, 11, f'=IF(B{rr}=0,"N/A",E{rr}/B{rr})')
        ws.cell(rr, 12, f'=IF(D{rr}=0,"N/A",E{rr}/D{rr})')
        ws.cell(rr, 13, f'=IF(C{rr}=0,"N/A",D{rr}/C{rr})')
        ws.cell(rr, 14, f'=ROUNDUP(G{rr}/$H$2,0)')
        ws.cell(rr, 15, f'=E{rr}/$H$2')
        ws.cell(rr, 16, 1 if pd.Timestamp(row.date).strftime("%Y-%m-%d") in outlier_set else 0)
    last = first + len(daily) - 1
    if len(daily) == 0:
        last = first - 1
    tot = last + 1
    avg = last + 2
    avg2 = last + 3
    n = len(daily)
    # 합계 행: order/sku는 기간 고유 수(값), 나머지 SUM, 비율은 합계 기준 수식
    ws.cell(tot, 1, "합계(기간)").font = BOLD
    ws.cell(tot, 2, int(r.period["order"])); ws.cell(tot, 3, int(r.period["sku"]))
    for col in (4, 5, 6, 7):
        L = get_column_letter(col)
        ws.cell(tot, col, f"=SUM({L}{first}:{L}{last})" if n else 0)
    ws.cell(tot, 8, f'=IF(G{tot}=0,"N/A",D{tot}/G{tot})')
    ws.cell(tot, 9, f'=IF(G{tot}=0,"N/A",E{tot}/G{tot})')
    ws.cell(tot, 10, f'=IF(B{tot}=0,"N/A",D{tot}/B{tot})')
    ws.cell(tot, 11, f'=IF(B{tot}=0,"N/A",E{tot}/B{tot})')
    ws.cell(tot, 12, f'=IF(D{tot}=0,"N/A",E{tot}/D{tot})')
    ws.cell(tot, 13, f'=IF(C{tot}=0,"N/A",D{tot}/C{tot})')
    ws.cell(tot, 14, f'=IF({n}=0,"N/A",ROUNDUP(G{tot}/({n}*$H$2),0))')   # 기간 tote/h = 총토트/(출고일수×시간)
    ws.cell(tot, 15, f'=IF({n}=0,"N/A",E{tot}/({n}*$H$2))')
    ws.cell(tot, 16, f"=COUNTIF(P{first}:P{last},1)" if n else 0)   # 특이일 수
    # AVG 행: 0·빈 값 제외 평균
    ws.cell(avg, 1, "AVG(일평균, 0 제외)").font = BOLD
    for col in range(2, 16):
        L = get_column_letter(col)
        ws.cell(avg, col, f'=IFERROR(AVERAGEIF({L}{first}:{L}{last},"<>0"),"N/A")' if n else "N/A")
    # AVG(특이일 제외) 행: P열(특이일=1)을 뺀 날짜의 0 제외 평균. 합계·기간 비율은 바뀌지 않는다.
    ws.cell(avg2, 1, "AVG(특이일 제외, 0 제외)").font = BOLD
    for col in range(2, 16):
        L = get_column_letter(col)
        ws.cell(avg2, col, f'=IFERROR(AVERAGEIFS({L}{first}:{L}{last},{L}{first}:{L}{last},"<>0",$P${first}:$P${last},0),"N/A")' if n else "N/A")
    for col in range(1, 17):
        ws.cell(tot, col).fill = SUM_FILL
        ws.cell(avg, col).fill = SUM_FILL
        ws.cell(avg2, col).fill = SUM_FILL
    for rr in range(first, avg2 + 1):
        for col in range(8, 16):
            ws.cell(rr, col).number_format = "0.00"
    ws.freeze_panes = ws.cell(first, 2)
    _autosize(ws, {"A": 20})


def write_compare_sheet(wb: Workbook, results: list[RunResult], baseline_index: int, sheet_names: list[str]) -> None:
    ws = wb.create_sheet("비교", 0)
    ws["A1"] = "시나리오 비교 (기간 합계 기준, 기준 시나리오 대비 증감)"
    ws["A1"].font = BOLD
    headers = ["시나리오", "시트", "DAS 셀", "토트 용량", "주문 단위", "작업시간", "기준",
               "order", "sku", "order line", "pcs", "tote",
               "OL/TOTE", "PCS/TOTE", "tote/h",
               "토트 증감", "토트 증감률(%)", "OL/TOTE 증감", "PCS/TOTE 증감", "tote/h 증감",
               "OL/TOTE(일평균)", "PCS/TOTE(일평균)", "토트 순위",
               "OL/TOTE(일평균·특이일 제외)", "PCS/TOTE(일평균·특이일 제외)", "특이일 수"]
    HR = 3
    _write_header(ws, HR, headers)
    first = HR + 1
    base_row = first + baseline_index
    cmp_df = compare(results, baseline_index)
    for i, r in enumerate(results):
        rr = first + i
        sn = sheet_names[i]
        p = r.period
        ws.cell(rr, 1, r.scenario.name); ws.cell(rr, 2, sn)
        ws.cell(rr, 1).data_type = "s"
        ws.cell(rr, 3, r.scenario.cells); ws.cell(rr, 4, r.scenario.tote_capacity)
        ws.cell(rr, 5, r.scenario.order_unit); ws.cell(rr, 6, r.scenario.hours_per_day)
        ws.cell(rr, 7, "●" if i == baseline_index else "")
        ws.cell(rr, 8, int(p["order"])); ws.cell(rr, 9, int(p["sku"]))
        # 시트 합계 행 참조: 시나리오 시트의 합계 행 위치 = 5 + n_days
        tot_row = 5 + r.n_days
        q = f"'{sn}'!"
        ws.cell(rr, 10, f"={q}D{tot_row}"); ws.cell(rr, 11, f"={q}E{tot_row}"); ws.cell(rr, 12, f"={q}G{tot_row}")
        ws.cell(rr, 13, f'=IF(L{rr}=0,"N/A",J{rr}/L{rr})')
        ws.cell(rr, 14, f'=IF(L{rr}=0,"N/A",K{rr}/L{rr})')
        ws.cell(rr, 15, f"={q}N{tot_row}")
        ws.cell(rr, 16, f"=L{rr}-L${base_row}")
        ws.cell(rr, 17, f'=IF(L${base_row}=0,"N/A",(L{rr}-L${base_row})/L${base_row}*100)')
        ws.cell(rr, 18, f'=IFERROR(M{rr}-M${base_row},"N/A")')
        ws.cell(rr, 19, f'=IFERROR(N{rr}-N${base_row},"N/A")')
        ws.cell(rr, 20, f'=IFERROR(O{rr}-O${base_row},"N/A")')
        ws.cell(rr, 21, r.avg.get("ol_per_tote")); ws.cell(rr, 22, r.avg.get("pcs_per_tote"))
        ws.cell(rr, 23, int(cmp_df.loc[i, "토트 순위"]))
        ws.cell(rr, 24, r.avg_ex.get("ol_per_tote")); ws.cell(rr, 25, r.avg_ex.get("pcs_per_tote"))
        ws.cell(rr, 26, len(r.flags.get("outlier_days", [])))
        for col in (13, 14, 17, 18, 19, 21, 22, 24, 25):
            ws.cell(rr, col).number_format = "0.00"
    ws.freeze_panes = ws.cell(first, 2)
    _autosize(ws, {"A": 22})
    _add_compare_charts(ws, first, first + len(results) - 1)


def _add_compare_charts(ws, first: int, last: int) -> None:
    """비교 표 아래에 엑셀 기본 차트 4개: 총 토트 / 히트율 / 토트 증감률 / tote·h.

    데이터는 비교 시트의 수식 셀을 참조한다 (캐시값이 있어 열자마자 그려진다).
    열 위치: A 시나리오, L tote, M OL/TOTE, N PCS/TOTE, O tote/h, Q 토트 증감률(%).
    """
    if last < first:
        return
    cats = Reference(ws, min_col=1, min_row=first, max_row=last)
    top = last + 3
    ws.cell(top - 1, 1, "차트 (기간 합계 기준)").font = BOLD

    def bar(title: str, cols: list[int], y_title: str, number_format: str, anchor: str, labels: bool = True) -> None:
        ch = BarChart()
        ch.type = "col"
        ch.title = title
        ch.y_axis.title = y_title
        ch.y_axis.number_format = number_format
        ch.y_axis.majorGridlines = None
        ch.x_axis.delete = False            # openpyxl 3.1+: 축을 명시해야 범주·눈금 라벨이 보인다
        ch.y_axis.delete = False
        ch.height, ch.width = 7.5, 14
        for c in cols:
            ch.add_data(Reference(ws, min_col=c, min_row=first - 1, max_row=last), titles_from_data=True)
        ch.set_categories(cats)
        for s in ch.series:
            s.invertIfNegative = False          # 음수(증감률) 막대가 흰색으로 비는 현상 방지
        if len(cols) > 1:
            ch.legend.position = "r"
        else:
            ch.legend = None
        if labels:
            ch.dataLabels = DataLabelList()
            ch.dataLabels.showVal = True
            ch.dataLabels.showSerName = False
            ch.dataLabels.showCatName = False
            ch.dataLabels.showLegendKey = False
            ch.dataLabels.showPercent = False
            ch.dataLabels.numFmt = number_format
        ws.add_chart(ch, anchor)

    bar("총 토트", [12], "tote", "#,##0", f"A{top}")
    bar("히트율 (OL/TOTE · PCS/TOTE)", [13, 14], "비율", "0.00", f"J{top}")
    bar("기준 대비 토트 증감률(%)", [17], "%", "0.0", f"A{top + 16}")
    bar("시간당 토트 (tote/h)", [15], "tote/h", "#,##0", f"J{top + 16}")


def write_settings_sheet(wb: Workbook, results: list[RunResult], dataset: Optional[dict], baseline_index: int) -> None:
    ws = wb.create_sheet("설정")
    rows = [
        ("생성일시 (한국)", datetime.now(timezone(timedelta(hours=9))).strftime("%Y-%m-%d %H:%M:%S")),
        ("엔진 버전", ENGINE_VERSION),
        ("데이터셋", (dataset or {}).get("name", "")),
        ("원본 파일", (dataset or {}).get("file_name", "")),
        ("시트", (dataset or {}).get("sheet", "")),
        ("행 수", (dataset or {}).get("row_count", "")),
        ("기간", f"{(dataset or {}).get('date_min', '')} ~ {(dataset or {}).get('date_max', '')}"),
        ("열 매핑", str((dataset or {}).get("mapping", ""))),
        ("필터", results[0].filters.describe() if results else ""),
        ("기준 시나리오", results[baseline_index].scenario.name if results else ""),
        ("데이터셋 ID", (dataset or {}).get("dataset_id", "")),
        ("실행 ID", ", ".join((dataset or {}).get("run_ids", []))),
    ]
    for i, (k, v) in enumerate(rows, start=1):
        ws.cell(i, 1, k).font = BOLD
        ws.cell(i, 2, v)
        if isinstance(v, str):
            ws.cell(i, 2).data_type = "s"
    r0 = len(rows) + 2
    _write_header(ws, r0, ["시나리오", "DAS 셀", "토트 용량", "주문 단위", "작업시간", "출고일 수"])
    for i, r in enumerate(results, start=1):
        ws.cell(r0 + i, 1, r.scenario.name); ws.cell(r0 + i, 2, r.scenario.cells)
        ws.cell(r0 + i, 1).data_type = "s"
        ws.cell(r0 + i, 3, r.scenario.tote_capacity); ws.cell(r0 + i, 4, r.scenario.order_unit)
        ws.cell(r0 + i, 5, r.scenario.hours_per_day); ws.cell(r0 + i, 6, r.n_days)
    _autosize(ws, {"B": 60})


def write_definitions_sheet(wb: Workbook) -> None:
    ws = wb.create_sheet("계산정의")
    _write_header(ws, 1, ["항목", "정의"])
    defs = [
        ("주문(order)", "시나리오의 주문 단위 값 1개 = DAS 셀 1칸. 배송처차수 전체 / 배송처(공백 앞부분)"),
        ("웨이브", "날짜별로 고유 주문 키를 문자 오름차순 정렬해 DAS 셀 수만큼 차례로 묶고 1,2,3… 날짜가 바뀌면 1부터"),
        ("토트", "(날짜, 웨이브, 상품코드)별 PCS 합계를 토트 용량으로 나눠 올림: ROUNDUP(PCS 합 / 토트 용량)"),
        ("기간 히트율", "Σorder line ÷ Σtote (날짜별 히트율의 평균과 구분)"),
        ("AVG", "0·빈 값을 뺀 날짜별 값의 단순 평균"),
        ("특이일", "주문 1건인 날 또는 OL/TOTE<1인 날. 합계·기간 비율에는 항상 포함"),
        ("AVG(특이일 제외)", "특이일을 뺀 날짜의 0 제외 단순 평균. AVG 계산에만 적용"),
        ("분모 0", "비율의 분모(tote, order, order line, sku)가 0이면 N/A. AVG에서 제외"),
        ("기간 tote/h", "ROUNDUP(총 tote ÷ (출고일 수 × 일 작업시간))"),
    ] + [(label, formula) for _, label, formula in METRIC_DEFS]
    for i, (k, v) in enumerate(defs, start=2):
        ws.cell(i, 1, k); ws.cell(i, 2, v)
    _autosize(ws, {"B": 90})


def export_excel(results: list[RunResult], path: str | Path, dataset: Optional[dict] = None,
                 baseline_index: int = 0) -> Path:
    """선택한 시나리오의 날짜 결과 + 비교 + 설정 + 계산정의를 한 엑셀로."""
    if not results:
        raise ValueError("내보낼 결과가 없습니다.")
    wb = Workbook()
    wb.remove(wb.active)
    names = [_sheet_name(i, r) for i, r in enumerate(results)]
    # 중복 시트명 방지
    seen: dict[str, int] = {}
    for i, n in enumerate(names):
        if n in seen:
            seen[n] += 1
            names[i] = (n[:28] + f"_{seen[n]}")[:31]
        else:
            seen[n] = 0
    for n, r in zip(names, results):
        write_scenario_sheet(wb, n, r)
    write_compare_sheet(wb, results, baseline_index, names)
    write_settings_sheet(wb, results, dataset, baseline_index)
    write_definitions_sheet(wb)
    out = Path(path)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    _cache_formula_values(out, results, baseline_index)
    return out


def _cache_formula_values(path, results, baseline_index):
    """Store verified engine values beside formulas so previews work before Excel recalculates."""
    metric_columns = {8: "ol_per_tote", 9: "pcs_per_tote", 10: "ol_per_order", 11: "pcs_per_order",
                      12: "pcs_per_ol", 13: "ol_per_sku", 14: "tote_per_h", 15: "pcs_per_h"}
    all_columns = {2: "order", 3: "sku", 4: "order_line", 5: "pcs", 6: "waves", 7: "tote", **metric_columns}
    caches = {}
    for i, result in enumerate(results):
        values = {}
        for j, row in enumerate(result.daily.to_dict("records"), 5):
            for col, key in metric_columns.items():
                values[f"{get_column_letter(col)}{j}"] = row[key]
        total = 5 + result.n_days
        for col, key in all_columns.items():
            values[f"{get_column_letter(col)}{total}"] = result.period[key]
            average = result.avg.get(key)
            if key == "waves":
                average = result.daily.waves.replace(0, float("nan")).mean()
            values[f"{get_column_letter(col)}{total + 1}"] = average
            excluded = result.avg_ex.get(key)
            if key == "waves":
                keep = ~result.daily["date"].dt.strftime("%Y-%m-%d").isin(result.flags.get("outlier_days", []))
                excluded = result.daily.loc[keep, "waves"].replace(0, float("nan")).mean()
            values[f"{get_column_letter(col)}{total + 2}"] = excluded
        values[f"P{total}"] = len(result.flags.get("outlier_days", []))
        caches[f"xl/worksheets/sheet{i + 2}.xml"] = values
    comparison = compare(results, baseline_index)
    values = {}
    for i, result in enumerate(results):
        row = i + 4
        p, c = result.period, comparison.iloc[i]
        for col, value in {10: p["order_line"], 11: p["pcs"], 12: p["tote"],
                           13: p["ol_per_tote"], 14: p["pcs_per_tote"], 15: p["tote_per_h"],
                           16: c["토트 증감"], 17: c["토트 증감률(%)"], 18: c["OL/TOTE 증감"],
                           19: c["PCS/TOTE 증감"], 20: c["tote/h 증감"]}.items():
            values[f"{get_column_letter(col)}{row}"] = value
    caches["xl/worksheets/sheet1.xml"] = values
    ns = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
    target = io.BytesIO()
    with zipfile.ZipFile(path) as source, zipfile.ZipFile(target, "w", zipfile.ZIP_DEFLATED) as output:
        for info in source.infolist():
            data = source.read(info.filename)
            if info.filename in caches:
                root = ET.fromstring(data)
                for cell in root.iter(ns + "c"):
                    coordinate = cell.get("r")
                    if cell.find(ns + "f") is None or coordinate not in caches[info.filename]:
                        continue
                    value = caches[info.filename][coordinate]
                    cached = cell.find(ns + "v")
                    if cached is None:
                        cached = ET.SubElement(cell, ns + "v")
                    if value is None or not math.isfinite(float(value)):
                        cell.set("t", "str")
                        cached.text = "N/A"
                    else:
                        cell.set("t", "n")
                        cached.text = repr(float(value))
                data = ET.tostring(root, encoding="utf-8", xml_declaration=True)
            output.writestr(info, data)
    Path(path).write_bytes(target.getvalue())


# ---------------------------------------------------------------- 설명 문장
def explain_comparison(results: list[RunResult], baseline_index: int = 0) -> str:
    """도구 결과 필드만으로 만든 근거 설명 (LLM 없이도 동작)."""
    if not results:
        return "비교할 결과가 없습니다."
    base = results[baseline_index]
    cmp_df = compare(results, baseline_index)
    bp = base.period
    lines = [
        f"기준 '{base.scenario.name}' (DAS {base.scenario.cells}셀 · 토트 {base.scenario.tote_capacity}PCS · "
        f"{base.scenario.order_unit} · {base.scenario.hours_per_day:g}h), 적용 범위 {base.filters.describe()}, "
        f"출고일 {base.n_days}일.",
        f"기준 결과: 총 토트 {bp['tote']:,} / OL/TOTE {bp['ol_per_tote']:.2f} / PCS/TOTE {bp['pcs_per_tote']:.2f} / "
        f"tote/h {bp['tote_per_h']}." if bp["tote"] else "기준 결과: 토트 0 (비율 N/A).",
    ]
    for i, r in enumerate(results):
        if i == baseline_index:
            continue
        row = cmp_df.loc[i]
        pct = row["토트 증감률(%)"]
        pct_s = f"{pct:+.1f}%" if pct is not None and pd.notna(pct) else "N/A"
        ol = f"{bp['ol_per_tote']:.2f} → {r.period['ol_per_tote']:.2f}" if r.period["ol_per_tote"] is not None and bp["ol_per_tote"] is not None else "N/A"
        lines.append(
            f"- '{r.scenario.name}': 총 토트 {r.period['tote']:,} ({int(row['토트 증감']):+,}, {pct_s}), "
            f"OL/TOTE {ol}, PCS/TOTE {_format_ratio(r.period['pcs_per_tote'])}, tote/h {r.period['tote_per_h']}."
        )
    outlier_days = base.flags.get("outlier_days", [])
    if outlier_days:
        lines.append(
            f"특이일 {len(outlier_days)}일(주문 1건 {len(base.flags.get('single_order_days', []))}, "
            f"OL/TOTE<1 {len(base.flags.get('low_hit_days', []))})은 합계·기간 비율에 포함. "
            f"특이일 제외 일평균 OL/TOTE {_format_ratio(base.avg_ex.get('ol_per_tote'))} "
            f"(포함 {base.avg_ex.get('days_included')}일 / 제외 {base.avg_ex.get('days_excluded')}일)."
        )
    min_tote = cmp_df["총 토트"].min()
    winners = cmp_df.loc[cmp_df["총 토트"] == min_tote, "시나리오"].tolist()
    if len(winners) == 1:
        lines.append(f"총 토트가 가장 적은 설정: '{winners[0]}' ({min_tote:,}). 설비 비용·인원은 이 비교에 포함되지 않는다.")
    else:
        lines.append(f"총 토트 최소 동률: {', '.join(repr(w) for w in winners)} ({min_tote:,}).")
    return "\n".join(lines)


def _format_ratio(value):
    return "N/A" if value is None or pd.isna(value) else f"{value:.2f}"
